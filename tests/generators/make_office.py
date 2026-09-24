import json, os, random, subprocess
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_TAB_ALIGNMENT
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

SP = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(SP, '..', 'fixtures', 'tables')
TMP = os.path.join(SP, '..', '.output', 'office')
os.makedirs(TMP, exist_ok=True)
random.seed(11)
FIRST = ['Jean', 'Marie', 'Paul', 'Lucie', 'Hugo', 'Emma']; LAST = ['Dupont', 'Martin', 'Bernard', 'Petit', 'Durand']
CITY = ['Paris', 'Lyon', 'Nantes', 'Lille', 'Brest']
def rows(n, prefix='C'):
    r = [['Réf.', 'Client', 'Ville', 'Qté', 'Montant']]
    for i in range(n):
        r.append([f'{prefix}{300+i}', f'{random.choice(FIRST)} {random.choice(LAST)}', random.choice(CITY), str(random.randint(1, 99)), f'{random.randint(10, 9999)},{random.randint(10, 99)} €'])
    return r
def gt(name, cells):
    json.dump({'pages': [{'cells': [{'text': t, 'box': None} for t in cells], 'rules': [], 'lines': []}]}, open(os.path.join(OUT, name + '.json'), 'w'), ensure_ascii=False)

def docx_table(name, style, data, font_size=10, bold_header=True):
    doc = Document()
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    if style: t.style = style
    for i, r in enumerate(data):
        for j, v in enumerate(r):
            p = t.cell(i, j).paragraphs[0]; run = p.add_run(v); run.font.size = Pt(font_size); run.bold = bold_header and i == 0
    path = os.path.join(TMP, name + '.docx'); doc.save(path); return path

files = {}
d1 = rows(15, 'C'); files['c01_docx_table_grid'] = (docx_table('c01_docx_table_grid', 'Table Grid', d1), sum(d1, []))
d2 = rows(15, 'D'); files['c02_docx_no_borders'] = (docx_table('c02_docx_no_borders', None, d2), sum(d2, []))
d3 = rows(20, 'E'); files['c03_docx_small_font'] = (docx_table('c03_docx_small_font', 'Table Grid', d3, font_size=8), sum(d3, []))

# Paragraphes avec taquets de tabulation (libellé → valeur)
doc = Document(); cells = []
for k, v in [('Nom :', 'Dupont'), ('Prénom :', 'Marie'), ('Adresse :', '12 rue des Lilas'), ('Ville :', 'Lyon')]:
    p = doc.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Cm(6)); p.add_run(k).bold = True; p.add_run('\t' + v); cells += [k, v]
path = os.path.join(TMP, 'c04_docx_tabs.docx'); doc.save(path); files['c04_docx_tabs'] = (path, cells)

# Tableur avec quadrillage imprimé et bordures
for name, grid, borders in [('c05_xlsx_gridlines', True, False), ('c06_xlsx_borders_fill', False, True)]:
    wb = openpyxl.Workbook(); ws = wb.active; data = rows(25, 'X'); thin = Side(style='thin')
    for i, r in enumerate(data, 1):
        for j, v in enumerate(r, 1):
            cell = ws.cell(row=i, column=j, value=v)
            if i == 1: cell.font = Font(bold=True, color='FFFFFF' if borders else '000000'); cell.fill = PatternFill('solid', fgColor='1F3A5F') if borders else PatternFill()
            if borders: cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if j >= 4: cell.alignment = Alignment(horizontal='right')
    for col, w in zip('ABCDE', [8, 18, 12, 6, 14]): ws.column_dimensions[col].width = w
    ws.print_options.gridLines = grid
    path = os.path.join(TMP, name + '.xlsx'); wb.save(path); files[name] = (path, sum(data, []))

for name, (path, cells) in files.items():
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', OUT, path], check=True, capture_output=True, timeout=180)
    gt(name, cells)
    print(name, len(cells), 'cells', os.path.getsize(os.path.join(OUT, name + '.pdf')))
